from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from models import FeedbackResult


class ExcelWriter:

    HEADER = [
        "No",
        "No.SEP",
        "Tgl. Verifikasi",
        "Biaya Riil RS",
        "Biaya Diajukan",
        "Biaya Disetujui",
    ]

    def write(
        self,
        result: FeedbackResult,
        output_path: str
    ):

        workbook = Workbook()

        #
        # Sheet RI
        #
        sheet_ri = workbook.active
        sheet_ri.title = "RI"

        self.write_sheet(
            sheet_ri,
            result.ri.rows
        )

        #
        # Sheet RJ
        #
        sheet_rj = workbook.create_sheet("RJ")

        self.write_sheet(
            sheet_rj,
            result.rj.rows
        )

        workbook.save(output_path)

    def write_sheet(
        self,
        sheet,
        rows
    ):

        #
        # Header
        #
        sheet.append(self.HEADER)

        self.style_header(sheet)

        #
        # Data
        #
        for row in rows:

            sheet.append([
                row.no,
                row.no_sep,
                row.tgl_verifikasi,
                row.biaya_riil_rs,
                row.biaya_diajukan,
                row.biaya_disetujui,
            ])

        #
        # Styling
        #
        self.style_body(sheet)

        self.auto_width(sheet)

    def style_header(
        self,
        sheet
    ):

        fill = PatternFill(
            fill_type="solid",
            fgColor="4472C4"
        )

        font = Font(
            bold=True,
            color="FFFFFF"
        )

        border = Border(

            left=Side(style="thin"),

            right=Side(style="thin"),

            top=Side(style="thin"),

            bottom=Side(style="thin"),
        )

        for cell in sheet[1]:

            cell.fill = fill

            cell.font = font

            cell.border = border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    def style_body(
        self,
        sheet
    ):

        border = Border(

            left=Side(style="thin"),

            right=Side(style="thin"),

            top=Side(style="thin"),

            bottom=Side(style="thin"),
        )

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                cell.border = border

        #
        # Format angka
        #
        for col in [
            "D",
            "E",
            "F"
        ]:

            for cell in sheet[col][1:]:

                cell.number_format = '#,##0'

                cell.alignment = Alignment(
                    horizontal="right"
                )

        #
        # Tengah
        #
        for cell in sheet["A"]:

            cell.alignment = Alignment(
                horizontal="center"
            )

        for cell in sheet["C"]:

            cell.alignment = Alignment(
                horizontal="center"
            )

    def auto_width(
        self,
        sheet
    ):

        for column in sheet.columns:

            length = 0

            letter = get_column_letter(
                column[0].column
            )

            for cell in column:

                if cell.value is None:
                    continue

                value = str(cell.value)

                if len(value) > length:
                    length = len(value)

            sheet.column_dimensions[
                letter
            ].width = length + 3